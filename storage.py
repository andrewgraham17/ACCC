"""SQLite storage for the Arm Care Command Center.
Every function writes immediately, so the app autosaves on each change.
Throw-log tables are unchanged from the K-State Throw Log app; the Command Center adds
assessments, photos, corrective programs, force-plate tests and the exercise library."""
from __future__ import annotations

import os
import re
import sqlite3
import time
import json
import unicodedata
import uuid
from datetime import date, datetime

import pandas as pd

DB_PATH = os.environ.get("THROW_DB", os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "command_center.db"))
HERE = os.path.dirname(os.path.abspath(__file__))

LOG_FIELDS = ["date", "session_type", "distance", "start_time", "end_time", "actual_throws",
              "velo_low", "velo_high", "warm_throws", "warm_velo_low", "warm_velo_high",
              "sixty_throws", "sixty_velo_low", "sixty_velo_high", "notes"]
NUM_FIELDS = {"distance", "actual_throws", "velo_low", "velo_high", "warm_throws", "warm_velo_low",
              "warm_velo_high", "sixty_throws", "sixty_velo_low", "sixty_velo_high"}

# 2026 pitching staff (jersey, name, position, year, bats/throws)
DEFAULT_ROSTER = [
    ("30", "Max Bettis", "RHP", "So.", "R/R"), ("33", "Sam Bettis", "RHP", "So.", "R/R"),
    ("26", "Noah Gartner", "RHP", "So.", "R/R"), ("31", "Cooper Jesperson", "RHP", "So.", "R/R"),
    ("23", "Jonah Reich", "RHP", "5th", "R/R"), ("32", "Jackson Baker", "RHP", "Sr.", "R/R"),
    ("7", "Trenton Buckley", "RHP", "Jr.", "R/R"), ("17", "Stevie Doty", "RHP", "Jr.", "R/R"),
    ("99", "Colin Driffill", "RHP", "Fr.", "R/R"), ("37", "Jackson Hulcher", "RHP", "Sr.", "R/R"),
    ("35", "Griffin Lewis", "RHP", "Jr.", "R/R"), ("50", "Max Nantais-Vlahovich", "RHP", "Fr.", "R/R"),
    ("28", "Joseph Pereira", "RHP", "Fr.", "L/R"), ("18", "Stocton Timbrook", "RHP", "Jr.", "R/R"),
    ("15", "Grady Westphal", "RHP", "So.", "R/R"), ("4", "Cordell Clinkingbeard", "RHP/INF", "Fr.", "R/R"),
    ("16", "Kyler Horsman", "RHP/OF", "Sr.", "R/R"), ("9", "Donte Lewis", "UTL", "Jr.", "R/R"),
    ("41", "Adam Arther", "LHP", "Sr.", "L/L"), ("0", "Kaden Dean", "LHP", "Fr.", "L/L"),
    ("25", "Grant Meert", "LHP", "Sr.", "L/L"),
]


# Same athlete, different spelling between the throwing workbook and the Command Center
ALIASES = {"joseph-pereira": "joe-pereira"}


def slugify(name: str) -> str:
    s = unicodedata.normalize("NFKD", name).lower()
    s = re.sub(r"[^\w\s-]", "", s).strip()
    s = re.sub(r"-+", "-", re.sub(r"[\s_]+", "-", s))
    return ALIASES.get(s, s)


_READY = set()
PLAYER_EXTRA_COLS = {"grp": "TEXT DEFAULT 'pitcher'", "throws_hand": "TEXT", "assess": "TEXT", "photo": "BLOB",
                     "corrective": "TEXT", "corrective_history": "TEXT", "created_at": "TEXT"}


def _db():
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    c = sqlite3.connect(DB_PATH)
    if DB_PATH in _READY:
        return c
    c.execute("""CREATE TABLE IF NOT EXISTS players(
        id TEXT PRIMARY KEY, name TEXT, number TEXT, position TEXT, year TEXT, bt TEXT,
        avg_top_velo REAL, peak_velo REAL, active INTEGER DEFAULT 1, sort INTEGER,
        target_type TEXT DEFAULT 'Game', target_date TEXT, planned_pitches REAL, planned_innings REAL)""")
    have = {r[1] for r in c.execute("PRAGMA table_info(players)")}
    for col, typ in PLAYER_EXTRA_COLS.items():      # upgrade an older throwing.db in place
        if col not in have:
            c.execute(f"ALTER TABLE players ADD COLUMN {col} {typ}")
    cols = ", ".join(f + (" REAL" if f in NUM_FIELDS else " TEXT") for f in LOG_FIELDS)
    c.execute(f"CREATE TABLE IF NOT EXISTS throws(id TEXT PRIMARY KEY, player TEXT, created REAL, updated REAL, {cols})")
    c.execute("CREATE TABLE IF NOT EXISTS overrides(player TEXT, day TEXT, day_type TEXT, PRIMARY KEY(player, day))")
    c.execute("CREATE TABLE IF NOT EXISTS force_tests(id TEXT PRIMARY KEY, player TEXT, date TEXT, data TEXT)")
    c.execute("CREATE TABLE IF NOT EXISTS exercise_library(id TEXT PRIMARY KEY, name TEXT, video_url TEXT)")
    c.execute("CREATE TABLE IF NOT EXISTS files(id TEXT PRIMARY KEY, player TEXT, kind TEXT, name TEXT, mime TEXT, data BLOB, created REAL)")
    c.commit()
    _READY.add(DB_PATH)
    return c


def _clean(field, v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if field == "date":
        if isinstance(v, datetime):
            v = v.date()
        return v.isoformat() if isinstance(v, date) else (str(v)[:10] or None)
    if field in ("start_time", "end_time"):
        return v.strftime("%H:%M") if hasattr(v, "strftime") else (str(v).strip()[:5] or None)
    if field in NUM_FIELDS:
        try:
            return float(v)
        except (TypeError, ValueError):
            return None
    s = str(v).strip()
    return s or None


# ───────── players ─────────
def load_players(include_inactive=False) -> pd.DataFrame:
    with _db() as c:
        df = pd.read_sql("SELECT * FROM players" + ("" if include_inactive else " WHERE active=1") +
                         " ORDER BY name", c)
    df["has_photo"] = df["photo"].notna()
    df = df.drop(columns=["photo"])
    df["target_date"] = [date.fromisoformat(v) if isinstance(v, str) and v else None for v in df["target_date"]]
    df["target_date"] = df["target_date"].astype(object)
    return df


def get_player(pid) -> dict | None:
    df = load_players(include_inactive=True)
    hit = df[df["id"] == pid]
    if hit.empty:
        return None
    r = hit.iloc[0].to_dict()
    for k in ("avg_top_velo", "peak_velo", "planned_pitches", "planned_innings"):
        if pd.isna(r.get(k)):
            r[k] = None
    if not isinstance(r.get("target_date"), date) or pd.isna(r.get("target_date")):
        r["target_date"] = None
    return r


def add_player(name, number="", position="RHP", year=None, bt=None, avg_top=None, peak=None, grp="pitcher", pid=None) -> str:
    pid = pid or slugify(name)
    with _db() as c:
        n = c.execute("SELECT COALESCE(MAX(sort), 0) + 1 FROM players").fetchone()[0]
        exists = c.execute("SELECT 1 FROM players WHERE id=?", (pid,)).fetchone()
        if exists:  # re-adding a removed player brings his log back
            c.execute("UPDATE players SET active=1, name=?, number=?, position=?, grp=? WHERE id=?", (name.strip(), number, position, grp, pid))
            if avg_top is not None:
                c.execute("UPDATE players SET avg_top_velo=? WHERE id=?", (avg_top, pid))
        else:
            c.execute("""INSERT INTO players(id, name, number, position, year, bt, avg_top_velo, peak_velo, active, sort, grp,
                         assess, corrective, corrective_history, created_at)
                         VALUES (?,?,?,?,?,?,?,?,1,?,?,'{}','{"exercises": [], "notes": ""}','[]',?)""",
                      (pid, name.strip(), number, position, year, bt, avg_top, peak, n, grp, datetime.utcnow().isoformat()))
    return pid


def update_player(pid, **fields):
    allowed = {"name", "number", "position", "year", "bt", "avg_top_velo", "peak_velo", "active",
               "target_type", "target_date", "planned_pitches", "planned_innings", "grp", "throws_hand", "photo"}
    with _db() as c:
        for k, v in fields.items():
            if k not in allowed:
                continue
            if isinstance(v, date):
                v = v.isoformat()
            if isinstance(v, float) and pd.isna(v):
                v = None
            c.execute(f"UPDATE players SET {k}=? WHERE id=?", (v, pid))


def remove_player(pid):
    """Hides him from the roster. His file and throw log are kept and come back if he's added again."""
    update_player(pid, active=0)


# ───────── throw log (row level) ─────────
def load_log(pid) -> pd.DataFrame:
    with _db() as c:
        df = pd.read_sql(f"SELECT id, created, {', '.join(LOG_FIELDS)} FROM throws WHERE player=?", c, params=(pid,))
    df = df.sort_values(["date", "created"], ascending=False, na_position="last").reset_index(drop=True)
    return df


def get_row(row_id) -> dict | None:
    with _db() as c:
        c.row_factory = sqlite3.Row
        r = c.execute("SELECT * FROM throws WHERE id=?", (row_id,)).fetchone()
    return dict(r) if r else None


def add_row(pid, rec: dict) -> str:
    rid = uuid.uuid4().hex[:12]
    now = time.time()
    with _db() as c:
        c.execute(f"INSERT INTO throws(id, player, created, updated, {', '.join(LOG_FIELDS)}) VALUES (?,?,?,?{',?' * len(LOG_FIELDS)})",
                  [rid, pid, rec.get("created", now), now] + [_clean(f, rec.get(f)) for f in LOG_FIELDS])
    return rid


def update_row(row_id, rec: dict):
    fields = [f for f in LOG_FIELDS if f in rec]
    if not fields:
        return
    with _db() as c:
        c.execute(f"UPDATE throws SET {', '.join(f + '=?' for f in fields)}, updated=? WHERE id=?",
                  [_clean(f, rec[f]) for f in fields] + [time.time(), row_id])


def delete_row(row_id):
    with _db() as c:
        c.execute("DELETE FROM throws WHERE id=?", (row_id,))


# ───────── overrides ─────────
def load_overrides(pid) -> dict:
    with _db() as c:
        return {date.fromisoformat(d): t for d, t in c.execute("SELECT day, day_type FROM overrides WHERE player=?", (pid,))}


def save_override(pid, day: date, day_type):
    with _db() as c:
        if not day_type or day_type == "(none)":
            c.execute("DELETE FROM overrides WHERE player=? AND day=?", (pid, day.isoformat()))
        else:
            c.execute("INSERT OR REPLACE INTO overrides VALUES (?,?,?)", (pid, day.isoformat(), day_type))


def clear_overrides(pid):
    with _db() as c:
        c.execute("DELETE FROM overrides WHERE player=?", (pid,))


# ───────── first run / import ─────────
def seed_if_empty(workbook_path=None, seed_json=None):
    """First run: load the Command Center (36 athletes, assessments, photos, correctives,
    force-plate tests, exercise library), then the throwing workbook (velo, throw log, planner)."""
    with _db() as c:
        if c.execute("SELECT COUNT(*) FROM players").fetchone()[0]:
            return False
    if seed_json and os.path.exists(seed_json):
        import_command_center(seed_json)
    for num, name, pos, yr, bt in DEFAULT_ROSTER:          # jersey / year / B-T from the throw-log roster
        pid = slugify(name)
        if get_player(pid) is None:
            add_player(name, num, pos, yr, bt)
        else:
            update_player(pid, number=num, position=pos, year=yr, bt=bt)
    if workbook_path and os.path.exists(workbook_path):
        import_workbook(workbook_path)
    return True


def import_command_center(path):
    """Load the JSON export of the Arm Care Command Center (players, forceTests, exerciseLibrary, throwLogs)."""
    with open(path, encoding="utf-8") as fh:
        d = json.load(fh)
    photo_dir = os.path.join(os.path.dirname(path), "photos")
    file_dir = os.path.join(os.path.dirname(path), "files")
    assess_keys = ["domER", "nonER", "domIR", "nonIR", "domArc", "nonArc", "domFlex", "nonFlex", "domElbowExt",
                   "nonElbowExt", "domHipExt", "nonHipExt", "erDiff", "irDiff", "flexDiff", "elbowExtDiff", "domGrip",
                   "nonGrip", "pinch", "oh", "trunk", "sls", "idb", "tspineRotFail", "tspineExtFail", "note", "sourceBatch"]
    with _db() as c:
        for pid, p in d.get("players", {}).items():
            corr = p.get("corrective") or {}
            exercises = corr.get("exercises")
            if not isinstance(exercises, list):
                exercises = [e for ph in corr.get("phases") or [] for e in (ph or {}).get("exercises", [])]
            corr_out = {"exercises": [{k: e.get(k, "") for k in ("name", "sets", "reps", "cues", "videoUrl")} for e in exercises],
                        "notes": corr.get("notes") or "", "startedAt": corr.get("startedAt")}
            photo = None
            pp = os.path.join(photo_dir, f"{pid}.jpg")
            if os.path.exists(pp):
                with open(pp, "rb") as fh:
                    photo = fh.read()
            c.execute("""INSERT OR REPLACE INTO players(id, name, number, position, avg_top_velo, active, sort, grp, throws_hand,
                         assess, photo, corrective, corrective_history, created_at, target_type)
                         VALUES (?,?,?,?,?,1,0,?,?,?,?,?,?,?,'Game')""",
                      (pid, p.get("name"), "", p.get("pos") or "", p.get("avgTopVelo"),
                       "pitcher" if p.get("group") == "pitcher" else "position", p.get("throw") or "",
                       json.dumps({k: p.get(k) for k in assess_keys}), photo, json.dumps(corr_out),
                       json.dumps(p.get("correctiveHistory") or []), p.get("createdAt")))
            pf = corr.get("programFile")
            if pf and pf.get("name"):
                fp = os.path.join(file_dir, f"{pid}__{pf['name']}")
                if os.path.exists(fp):
                    with open(fp, "rb") as fh:
                        c.execute("INSERT OR REPLACE INTO files VALUES (?,?,?,?,?,?,?)",
                                  (f"{pid}-program", pid, "program", pf["name"], pf.get("type") or "", fh.read(), time.time()))
        for tid, t in d.get("forceTests", {}).items():
            c.execute("INSERT OR REPLACE INTO force_tests VALUES (?,?,?,?)", (tid, t.get("playerId"), t.get("date"), json.dumps(t)))
        for lid, e in d.get("exerciseLibrary", {}).items():
            c.execute("INSERT OR REPLACE INTO exercise_library VALUES (?,?,?)", (lid, e.get("name", ""), e.get("videoUrl", "")))
    # Throw-log entries saved in the Command Center (duration-based). The workbook import that follows
    # replaces any pitcher it has rows for, so nothing is ever doubled.
    for pid, entries in d.get("throwLogs", {}).items():
        for e in entries.values():
            dur = e.get("durationMin")
            add_row(pid, {"date": e.get("date"), "session_type": e.get("sessionType"), "distance": e.get("distanceFt"),
                          "start_time": "12:00" if dur else None,
                          "end_time": f"{12 + int(dur) // 60:02d}:{int(dur) % 60:02d}" if dur else None,
                          "actual_throws": e.get("actualThrowCount"), "velo_low": e.get("veloLow"), "velo_high": e.get("veloHigh"),
                          "warm_throws": e.get("warmupThrows"), "warm_velo_low": e.get("warmupVeloLow"),
                          "warm_velo_high": e.get("warmupVeloHigh"), "sixty_throws": e.get("sixtyThrows"),
                          "sixty_velo_low": e.get("sixtyVeloLow"), "sixty_velo_high": e.get("sixtyVeloHigh"),
                          "notes": e.get("notes")})


def import_workbook(src) -> tuple[int, int]:
    """Import Avg Top Velo / Peak Velo (Roster tab), the Throw Log and planner inputs from the ACWR workbook."""
    import warnings
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(src)
    n_players = 0
    for row in wb["Roster"].iter_rows(min_row=5, values_only=True):
        if row[0]:
            name = str(row[0]).strip()
            pid = slugify(name)
            if get_player(pid) is None:
                add_player(name, position=row[1] or "RHP")
            update_player(pid, avg_top_velo=row[3], peak_velo=row[2])
            n_players += 1
    colmap = {1: "date", 3: "session_type", 4: "distance", 5: "start_time", 6: "end_time", 8: "actual_throws",
              10: "velo_low", 11: "velo_high", 18: "warm_throws", 19: "warm_velo_low", 20: "warm_velo_high",
              24: "sixty_throws", 25: "sixty_velo_low", 26: "sixty_velo_high", 30: "notes"}
    ws = wb["Throw Log"]
    by_player: dict = {}
    for r in range(5, ws.max_row + 1):
        name = ws.cell(r, 2).value
        if not name:
            continue
        rec = {}
        for col, k in colmap.items():
            v = ws.cell(r, col).value
            rec[k] = None if isinstance(v, str) and v.startswith("=") else v
        rec["created"] = 1000 + r
        by_player.setdefault(slugify(str(name)), []).append(rec)
    n = 0
    for pid, recs in by_player.items():
        if get_player(pid) is None:
            add_player(pid.replace("-", " ").title())
        with _db() as c:
            c.execute("DELETE FROM throws WHERE player=?", (pid,))   # replace, never duplicate
        for rec in recs:
            add_row(pid, rec)
            n += 1
    if "Next Outing Planner" in wb.sheetnames:
        nop = wb["Next Outing Planner"]
        if nop["B4"].value and isinstance(nop["B6"].value, datetime):
            update_player(slugify(str(nop["B4"].value)), target_type=nop["B5"].value or "Game",
                          target_date=nop["B6"].value.date(), planned_pitches=nop["B23"].value,
                          planned_innings=nop["D23"].value)
    return n_players, n


# ───────── assessment / photo / correctives ─────────
def _json_col(pid, col, default):
    with _db() as c:
        r = c.execute(f"SELECT {col} FROM players WHERE id=?", (pid,)).fetchone()
    try:
        return json.loads(r[0]) if r and r[0] else default
    except (TypeError, ValueError):
        return default


def get_assess(pid) -> dict:
    return _json_col(pid, "assess", {})


def all_assess() -> dict:
    with _db() as c:
        return {pid: (json.loads(a) if a else {}) for pid, a in c.execute("SELECT id, assess FROM players WHERE active=1")}


def update_assess(pid, **fields):
    a = get_assess(pid)
    a.update(fields)
    with _db() as c:
        c.execute("UPDATE players SET assess=? WHERE id=?", (json.dumps(a), pid))


def get_photo(pid):
    with _db() as c:
        r = c.execute("SELECT photo FROM players WHERE id=?", (pid,)).fetchone()
    return r[0] if r else None


def get_corrective(pid) -> dict:
    d = _json_col(pid, "corrective", {})
    d.setdefault("exercises", [])
    d.setdefault("notes", "")
    return d


def save_corrective(pid, corrective: dict):
    with _db() as c:
        c.execute("UPDATE players SET corrective=? WHERE id=?", (json.dumps(corrective), pid))


def get_corrective_history(pid) -> list:
    return _json_col(pid, "corrective_history", [])


def start_new_program(pid):
    """Archive the current program into history and start a blank one."""
    cur, hist = get_corrective(pid), get_corrective_history(pid)
    if cur.get("exercises") or cur.get("notes"):
        hist.insert(0, dict(cur, archivedAt=datetime.utcnow().isoformat()))
    with _db() as c:
        c.execute("UPDATE players SET corrective=?, corrective_history=? WHERE id=?",
                  (json.dumps({"exercises": [], "notes": "", "startedAt": datetime.utcnow().isoformat()}), json.dumps(hist), pid))


# ───────── files (program PDFs / videos, throwing-log workbooks) ─────────
def get_file(pid, kind):
    with _db() as c:
        r = c.execute("SELECT name, mime, data FROM files WHERE player=? AND kind=?", (pid, kind)).fetchone()
    return dict(name=r[0], mime=r[1], data=r[2]) if r else None


def save_file(pid, kind, name, mime, data: bytes):
    with _db() as c:
        c.execute("INSERT OR REPLACE INTO files VALUES (?,?,?,?,?,?,?)", (f"{pid}-{kind}", pid, kind, name, mime, data, time.time()))


def delete_file(pid, kind):
    with _db() as c:
        c.execute("DELETE FROM files WHERE player=? AND kind=?", (pid, kind))


# ───────── force plate (ForceDecks CMJ) ─────────
def load_force() -> pd.DataFrame:
    with _db() as c:
        rows = [dict(json.loads(d), id=i) for i, d in c.execute("SELECT id, data FROM force_tests")]
    df = pd.DataFrame(rows)
    if df.empty:
        return df
    if "excluded" not in df:
        df["excluded"] = False
    df["excluded"] = df["excluded"].fillna(False).astype(bool)
    return df


def import_force_csv(src, players_df) -> tuple[int, list]:
    """Import a ForceDecks CMJ export. Matches 'Name' to roster names; returns (#imported, unmatched names)."""
    raw = pd.read_csv(src)
    cols = {c.lower().strip(): c for c in raw.columns}

    def col(*names):
        for n in names:
            for k, v in cols.items():
                if n in k:
                    return v
        return None
    c_name, c_date = col("name"), col("date")
    pick = {"jumpHeightCm": col("jump height (imp-mom)", "jump height"), "peakPowerWkg": col("peak power / bm"),
            "rsiMod": col("rsi-modified", "rsi modified"), "concImpulseNs": col("concentric impulse"),
            "eccAsymPct": col("eccentric peak force % (asym)", "eccentric peak force asym"),
            "cmDepthCm": col("countermovement depth"), "bwKg": col("bw [kg]", "body weight", "bodyweight")}
    if not c_name or not c_date:
        raise ValueError("The CSV needs a Name column and a Date column (ForceDecks test export).")
    by_name = {n.lower(): i for i, n in zip(players_df["id"], players_df["name"])}
    n, unmatched = 0, set()
    with _db() as c:
        for _, r in raw.iterrows():
            name = str(r[c_name]).strip()
            pid = by_name.get(name.lower()) or (slugify(name) if slugify(name) in by_name.values() else None)
            d = pd.to_datetime(r[c_date], errors="coerce")
            if pd.isna(d):
                continue
            rec = {"date": d.date().isoformat(), "time": d.strftime("%H%M"), "exportName": name, "testType": "CMJ",
                   "playerId": pid or f"unmatched-{slugify(name)}", "matchedBy": "exact" if pid else "unmatched",
                   "source": getattr(src, "name", "upload")}
            for k, cn in pick.items():
                if cn is not None:
                    try:
                        val = str(r[cn]).strip()
                        side = val[-1] if val[-1:] in ("L", "R") else None
                        rec[k] = float(val.rstrip("LR ").strip())
                        if k == "eccAsymPct" and side:
                            rec["eccAsymSide"] = side
                    except (ValueError, IndexError):
                        rec[k] = None
            if rec.get("bwKg"):
                rec["bwLb"] = round(rec["bwKg"] * 2.20462, 1)
            if not pid:
                unmatched.add(name)
            tid = f"{rec['date']}_{rec['playerId']}_{rec['time']}"
            c.execute("INSERT OR REPLACE INTO force_tests VALUES (?,?,?,?)", (tid, rec["playerId"], rec["date"], json.dumps(rec)))
            n += 1
    return n, sorted(unmatched)


# ───────── exercise library ─────────
def load_library() -> pd.DataFrame:
    with _db() as c:
        return pd.read_sql("SELECT id, name, video_url FROM exercise_library ORDER BY name COLLATE NOCASE", c)


def save_library_item(name, video_url):
    if not (name or "").strip():
        return
    with _db() as c:
        c.execute("INSERT OR REPLACE INTO exercise_library VALUES (?,?,?)", (slugify(name), name.strip(), (video_url or "").strip()))


def delete_library_item(lid):
    with _db() as c:
        c.execute("DELETE FROM exercise_library WHERE id=?", (lid,))


def delete_player_forever(pid):
    with _db() as c:
        for t, col in (("players", "id"), ("throws", "player"), ("overrides", "player"), ("files", "player")):
            c.execute(f"DELETE FROM {t} WHERE {col}=?", (pid,))
